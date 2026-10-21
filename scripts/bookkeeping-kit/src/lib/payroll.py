# ── bookkeeping kit: payroll adapter (inlined into payroll nodes by gen-ext.js) ──
# Xero Payroll AU through xero-payroll-au (reads take xero_tenant_id, so each call lands in the client's own
# organisation) and Deputy through deputy. Read-only: the checks list corrections, the payroll officer makes them.
import csv
import io

from openpyxl import Workbook, load_workbook

BK_PAYROLL = "Bookkeeping/Payroll"
BK_PAYRUNS = "Bookkeeping/PayRuns"
BK_PAYROLL_FILES = "Bookkeeping Payroll"
PAYROLL = "xero-payroll-au"
DEPUTY = "deputy"
PAY_WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
PAY_BANDS = ["ordinary", "early", "late", "saturday", "sunday", "public_holiday", "ot_first", "ot_after"]
PAY_BAND_NAMES = {"ordinary": "ordinary", "early": "early (before start of day)", "late": "late (after end of day)",
                  "saturday": "Saturday", "sunday": "Sunday", "public_holiday": "public holiday",
                  "ot_first": "overtime (first hours)", "ot_after": "overtime (after that)"}


def pay_call(tool, tenant_id=None, **args):
    if tenant_id:
        args["xero_tenant_id"] = tenant_id
    exc = None  # bound first for the platform lint (see xero_call)
    try:
        return api.connectors.invoke(server=PAYROLL, tool=tool, args=args)
    except Exception as exc:
        raise RuntimeError("Xero Payroll: " + str(exc)[:500]) from None


def pay_rows(result, key):
    if isinstance(result, dict):
        rows = result.get(key)
        return rows if isinstance(rows, list) else []
    return result if isinstance(result, list) else []


def pay_date(value):
    """Xero Payroll dates arrive as '/Date(1759276800000+0000)/' or ISO; Deputy's as ISO or epoch seconds."""
    s = str(value or "")
    m = re.match(r"/Date\((-?\d+)", s)
    if m:
        return datetime.datetime.fromtimestamp(int(m.group(1)) / 1000, datetime.timezone.utc).date()
    try:
        return datetime.date.fromisoformat(s[:10])
    except ValueError:
        return None


def pay_xero_dt(d):
    return "DateTime(%d,%d,%d)" % (d.year, d.month, d.day)


def pay_period_end(start, kind):
    kind = str(kind or "").upper()
    if kind == "WEEKLY":
        return start + datetime.timedelta(days=6)
    if kind == "FOURWEEKLY":
        return start + datetime.timedelta(days=27)
    if kind in ("MONTHLY", "QUARTERLY"):
        months = 3 if kind == "QUARTERLY" else 1
        y, mo = start.year + (start.month - 1 + months) // 12, (start.month - 1 + months) % 12 + 1
        return datetime.date(y, mo, min(start.day, 28)) - datetime.timedelta(days=1)
    if kind == "TWICEMONTHLY":
        return start + datetime.timedelta(days=14)
    return start + datetime.timedelta(days=13)  # FORTNIGHTLY


def pay_calendar(tenant_id, name):
    """The client's pay calendar by its Xero name: {id, name, type, start, end, payment}, or None."""
    for c in pay_rows(pay_call("list_payroll_calendars", tenant_id), "PayrollCalendars"):
        if bk_norm(c.get("Name")) == bk_norm(name):
            start = pay_date(c.get("StartDate"))
            if not start:
                return None
            return {"id": c.get("PayrollCalendarID"), "name": c.get("Name"), "type": c.get("CalendarType"),
                    "start": str(start), "end": str(pay_period_end(start, c.get("CalendarType"))),
                    "payment": str(pay_date(c.get("PaymentDate")) or "")}
    return None


def pay_prep_date(payment, weekday):
    """The preparation weekday on or before the pay day."""
    target = PAY_WEEKDAYS.index(weekday) if weekday in PAY_WEEKDAYS else 0
    return payment - datetime.timedelta(days=(payment.weekday() - target) % 7)


def pay_name(e):
    return re.sub(r"\s+", " ", (str(e.get("FirstName") or "") + " " + str(e.get("LastName") or "")).strip())


def pay_employees(tenant_id, calendar_id=""):
    """Active employees (full records, one get_employee each), optionally only those on the pay calendar."""
    out = []
    for page in range(1, 11):
        rows = pay_rows(pay_call("list_employees", tenant_id, where='Status=="ACTIVE"', page=page), "Employees")
        for r in rows:
            if calendar_id and r.get("PayrollCalendarID") and r.get("PayrollCalendarID") != calendar_id:
                continue
            full = pay_rows(pay_call("get_employee", tenant_id, employeeId=r["EmployeeID"]), "Employees")
            out.append(full[0] if full else r)
        if len(rows) < 100:
            break
    return out


def pay_basis(e):
    basis = str(((e.get("TaxDeclaration") or {}).get("EmploymentBasis")) or e.get("EmploymentBasis") or "").upper()
    return {"FULLTIME": "full-time", "PARTTIME": "part-time", "CASUAL": "casual"}.get(basis, basis.lower() or "unknown")


def pay_items(tenant_id):
    items = (pay_call("list_pay_items", tenant_id) or {})
    items = items.get("PayItems", items) if isinstance(items, dict) else {}
    earnings = {r.get("EarningsRateID"): r for r in (items.get("EarningsRates") or []) if isinstance(r, dict)}
    leave = {r.get("LeaveTypeID"): r for r in (items.get("LeaveTypes") or []) if isinstance(r, dict)}
    return {"earnings": earnings, "leave": leave}


def pay_terminated(tenant_id, calendar_id, today):
    """Employees terminated in the last PAY_TERMINATED_DAYS (full records), optionally only those on the pay calendar.
    The list's own termination date (when it has one) saves fetching older leavers in full."""
    out = []
    for page in range(1, 11):
        rows = pay_rows(pay_call("list_employees", tenant_id, where='Status=="TERMINATED"', page=page), "Employees")
        for r in rows:
            if calendar_id and r.get("PayrollCalendarID") and r.get("PayrollCalendarID") != calendar_id:
                continue
            gone = pay_date(r.get("TerminationDate"))
            if gone and (today - gone).days > PAY_TERMINATED_DAYS:
                continue
            full = pay_rows(pay_call("get_employee", tenant_id, employeeId=r["EmployeeID"]), "Employees")
            out.append(full[0] if full else r)
        if len(rows) < 100:
            break
    return out


def pay_payslip(tenant_id, payslip_id):
    """A payslip in full (earnings, leave, super and reimbursement lines)."""
    result = pay_call("get_payslip", tenant_id, payslipId=payslip_id)
    slip = result.get("Payslip") if isinstance(result, dict) else None
    if isinstance(slip, list):
        slip = slip[0] if slip else None
    return slip if isinstance(slip, dict) else {}


def pay_timesheets(tenant_id, start, end):
    where = "StartDate>=%s AND EndDate<=%s" % (pay_xero_dt(start), pay_xero_dt(end))
    return pay_rows(pay_call("list_timesheets", tenant_id, where=where), "Timesheets")


def pay_draft_run(tenant_id, calendar_id, start):
    """The draft pay run for the period with its payslip summaries, or None."""
    where = 'PayrollCalendarID==Guid("%s") AND PayRunStatus=="DRAFT"' % calendar_id
    for r in pay_rows(pay_call("list_pay_runs", tenant_id, where=where), "PayRuns"):
        if pay_date(r.get("PayRunPeriodStartDate")) == start:
            full = pay_rows(pay_call("get_pay_run", tenant_id, payRunId=r["PayRunID"]), "PayRuns")
            return full[0] if full else r
    return None


def deputy_rows(tool, **args):
    out = []
    for page in range(20):
        rows = api.connectors.invoke(server=DEPUTY, tool=tool, args=dict(args, max=500, start=page * 500))
        rows = rows if isinstance(rows, list) else (rows or {}).get("data") or []
        out.extend(r for r in rows if isinstance(r, dict))
        if len(rows) < 500:
            break
    return out


def pay_band_of(name, mapping):
    """Which band a Xero earnings rate pays, from the pay guide's pay_items sheet or its name."""
    n = bk_norm(name)
    if n in mapping:
        return mapping[n]
    if "public holiday" in n or re.search(r"\bph\b", n):
        return "public_holiday"
    if "saturday" in n or re.search(r"\bsat\b", n):
        return "saturday"
    if "sunday" in n or re.search(r"\bsun\b", n):
        return "sunday"
    if "overtime" in n or re.search(r"\bot\b", n):
        return "ot_after" if ("after" in n or "thereafter" in n) else "ot_first"
    if "before 7" in n or "early" in n:
        return "early"
    if "after 6" in n or "late" in n or "evening" in n:
        return "late"
    return "ordinary"


def pay_guide(path):
    """The client's pay guide workbook as plain data, or None when there's no file. Sheets (any order, headers in row 1):
    rates (award, classification, basis, age_from, age_to, band, rate, effective_from); employees (employee, award,
    classification, state, special_rate, special_until, notes); rules (award, day_start, day_end, ot_after_hours,
    ot_first_hours); holidays (date, state, name); pay_items (pay_item, band)."""
    try:
        local = api.files.get(file=path)
    except Exception:
        return None
    book = load_workbook(local, read_only=True, data_only=True)
    guide = {"rates": [], "employees": {}, "rules": {}, "holidays": {}, "pay_items": {}}
    for ws in book.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if not rows:
            continue
        head = [bk_norm(h).replace(" ", "_") for h in rows[0]]
        recs = [dict(zip(head, r)) for r in rows[1:] if any(v not in (None, "") for v in r)]
        title = bk_norm(ws.title).replace(" ", "_")
        if title == "rates":
            for r in recs:
                band = str(r.get("band") or "ordinary").strip().lower().replace(" ", "_")
                guide["rates"].append({"award": bk_norm(r.get("award")), "classification": bk_norm(r.get("classification")),
                                       "basis": bk_norm(r.get("basis")) or "casual", "age_from": bk_money(r.get("age_from")),
                                       "age_to": bk_money(r.get("age_to")), "band": band, "rate": bk_money(r.get("rate")),
                                       "effective_from": str(pay_date(r.get("effective_from")) or "")})
        elif title == "employees":
            for r in recs:
                guide["employees"][bk_norm(r.get("employee"))] = {
                    "award": bk_norm(r.get("award")), "classification": bk_norm(r.get("classification")),
                    "state": str(r.get("state") or "").strip().upper(), "special_rate": bk_money(r.get("special_rate")),
                    "special_until": str(pay_date(r.get("special_until")) or ""), "notes": str(r.get("notes") or "")}
        elif title == "rules":
            for r in recs:
                guide["rules"][bk_norm(r.get("award"))] = {
                    "day_start": str(r.get("day_start") or "07:00")[:5], "day_end": str(r.get("day_end") or "18:00")[:5],
                    "ot_after_hours": bk_money(r.get("ot_after_hours")), "ot_first_hours": bk_money(r.get("ot_first_hours")) or 0}
        elif title == "holidays":
            for r in recs:
                d = pay_date(r.get("date"))
                for st in [s.strip().upper() for s in str(r.get("state") or "ALL").split(",") if s.strip()]:
                    if d:
                        guide["holidays"].setdefault(st, []).append(str(d))
        elif title == "pay_items":
            for r in recs:
                guide["pay_items"][bk_norm(r.get("pay_item"))] = str(r.get("band") or "ordinary").strip().lower().replace(" ", "_")
    return guide


def pay_rate(guide, award, classification, basis, age, band, on_date):
    """The pay guide's rate for this employee and band on a date, or None."""
    best = None
    for r in (guide or {}).get("rates") or []:
        if r["award"] != award or r["classification"] != classification or r["band"] != band or r["rate"] is None:
            continue
        if r["basis"] not in (basis, "any"):
            continue
        if age is not None and ((r["age_from"] is not None and age < r["age_from"]) or (r["age_to"] is not None and age > r["age_to"])):
            continue
        if r["effective_from"] and r["effective_from"] > str(on_date):
            continue
        if best is None or r["effective_from"] > best["effective_from"]:
            best = r
    return best["rate"] if best else None


def pay_hm(text, default):
    m = re.match(r"(\d{1,2}):(\d{2})", str(text or ""))
    return (int(m.group(1)), int(m.group(2))) if m else default


def pay_split(start_dt, end_dt, worked_hours, rule, holidays):
    """A shift's worked hours by band: public holiday, Saturday, Sunday, else early / ordinary / late by the award's
    day start and end; then daily overtime past the award's threshold. Breaks are spread pro rata."""
    span = (end_dt - start_dt).total_seconds() / 3600
    hours = worked_hours if worked_hours is not None else span
    bands = {b: 0.0 for b in PAY_BANDS}
    if span <= 0 or hours <= 0:
        return bands
    day = start_dt.date()
    if str(day) in holidays:
        bands["public_holiday"] = hours
    elif day.weekday() == 5:
        bands["saturday"] = hours
    elif day.weekday() == 6:
        bands["sunday"] = hours
    else:
        sh, sm = pay_hm(rule.get("day_start"), (7, 0))
        eh, em = pay_hm(rule.get("day_end"), (18, 0))
        day_start = start_dt.replace(hour=sh, minute=sm, second=0, microsecond=0)
        day_end = start_dt.replace(hour=eh, minute=em, second=0, microsecond=0)
        early = max(0.0, (min(end_dt, day_start) - start_dt).total_seconds() / 3600)
        late = max(0.0, (end_dt - max(start_dt, day_end)).total_seconds() / 3600)
        scale = hours / span
        bands["early"] = round(early * scale, 4)
        bands["late"] = round(late * scale, 4)
        bands["ordinary"] = round(hours - bands["early"] - bands["late"], 4)
        limit = rule.get("ot_after_hours")
        if limit and hours > limit:
            ot = hours - limit
            first = min(ot, rule.get("ot_first_hours") or 0)
            bands["ot_first"], bands["ot_after"] = round(first, 4), round(ot - first, 4)
            for b in ("ordinary", "late", "early"):
                take = min(bands[b], ot)
                bands[b] = round(bands[b] - take, 4)
                ot -= take
    return bands


def pay_table(path):
    """Rows of a spreadsheet or CSV in Knowledge as dicts keyed by normalised header (first row that has one)."""
    local = api.files.get(file=path)
    if path.lower().endswith(".csv"):
        with open(local, encoding="utf-8-sig", errors="replace") as fh:
            rows = list(csv.reader(io.StringIO(fh.read())))
    else:
        book = load_workbook(local, read_only=True, data_only=True)
        rows = [list(r) for r in book.worksheets[0].iter_rows(values_only=True)]
    for i, r in enumerate(rows[:20]):
        head = [bk_norm(h) for h in r]
        if any(h in ("employee", "name", "employee name", "staff") for h in head) and any(
                ("amount" in h or "total" in h) for h in head):
            return [dict(zip(head, x)) for x in rows[i + 1:] if any(v not in (None, "") for v in x)]
    return []


def pay_guide_template(dest, client_name):
    """Save a pay guide workbook with the sheets the checks read, and one EXAMPLE row each, to Knowledge."""
    book = Workbook()
    sheets = {
        "rates": (["award", "classification", "basis", "age_from", "age_to", "band", "rate", "effective_from"],
                  [["EXAMPLE GRIA", "Retail Employee Level 1", "casual", 21, "", "ordinary", 33.19, "2026-07-01"],
                   ["EXAMPLE GRIA", "Retail Employee Level 1", "casual", 21, "", "saturday", 39.83, "2026-07-01"],
                   ["EXAMPLE GRIA", "Retail Employee Level 1", "casual", 16, 16, "ordinary", 16.60, "2026-07-01"]]),
        "employees": (["employee", "award", "classification", "state", "special_rate", "special_until", "notes"],
                      [["EXAMPLE Jane Citizen", "EXAMPLE GRIA", "Retail Employee Level 1", "NSW", "", "", "Driver"]]),
        "rules": (["award", "day_start", "day_end", "ot_after_hours", "ot_first_hours"],
                  [["EXAMPLE GRIA", "07:00", "18:00", 9, 3]]),
        "holidays": (["date", "state", "name"], [["2026-12-25", "ALL", "Christmas Day"]]),
        "pay_items": (["pay_item", "band"], [["EXAMPLE Casual Sat", "saturday"]]),
    }
    blank = book.active
    for title, (head, rows) in sheets.items():
        ws = book.create_sheet(title=title)
        ws.append(head)
        for r in rows:
            ws.append(r)
    if blank is not None:
        book.remove(blank)
    local = api.files.scratch_path(dest.split("/")[-1])
    book.save(local)
    api.files.put(file=dest, scratch_file=local)
    return dest


def pay_run_load(todo_id="", client_name=""):
    """(path, run) for a pay run's WorkQ item (or its review / approval item), or a client's newest open pay run."""
    if todo_id:
        link = bk_read_json(BK_PAYRUNS + "/" + todo_id + ".json") or {}
        if link.get("link"):
            path = BK_PAYRUNS + "/" + link["link"] + ".json"
            return path, bk_read_json(path) or {}
        return "", {}
    if client_name:
        slug = bk_slug(client_name)
        runs = [f for f in api.files.list(path=BK_PAYRUNS)
                if str(f.get("name") or "").startswith(slug + "-") and str(f.get("name") or "").endswith(".json")]
        for f in sorted(runs, key=lambda f: str(f.get("name")), reverse=True):
            run = bk_read_json(f["path"]) or {}
            if run.get("period_start") and run.get("status") not in ("done", "cancelled"):
                return f["path"], run
    return "", {}


def pay_link(todo_id, run):
    """Point a WorkQ item at its pay run record, so its form submissions find the pay run."""
    bk_write_json(BK_PAYRUNS + "/" + todo_id + ".json", {"link": run["slug"] + "-" + run["period_start"]})


# ── Payroll checks from the mySMB checks catalogue (run by Check pay run) ──
# Privacy (the item owner's sign-off, 21 Oct 2026): findings name the employee and show pay figures (hours, rates,
# amounts) only. Tax file numbers, bank details, dates of birth, addresses and contact details are never shown: the checks
# about them state a fact ("no tax file number after 28 days", "paid into the same bank account as ..."), and bank
# accounts are compared in memory as hashes, never stored. Pay run items are visible to their creator (the automation
# owner) and assignees only.
PAY_CHECKS = {
    "CHK-147": ("Terminated Employee Is Missing an STP Cessation Reason", "High"),
    "CHK-148": ("Employee Has No Tax File Number After 28 Days", "Medium"),
    "CHK-149": ("Employee Has a Negative Leave Balance", "Low"),
    "CHK-150": ("Casual Employee Is Accruing Annual Leave", "Low"),
    "CHK-151": ("Reimbursements Exceed Wages", "Low"),
    "CHK-153": ("Employee Paid Before Their Start Date", "Low"),
    "CHK-155": ("Terminated Employee Still Being Paid", "High"),
    "CHK-156": ("Employee Paid Below The Wage Floor", "High"),
    "CHK-157": ("Super Liability Not Paid", "High"),
    "CHK-159": ("Employee Details Missing After Pay Run", "Medium"),
    "CHK-161": ("Employee Leave Entitlements Missing", "Medium"),
    "CHK-162": ("Employee Bank Accounts Same", "Low"),
    "CHK-163": ("Employee Birthday (Junior Employee)", "Low"),
    "CHK-164": ("Check Employee Super Contribution", "Low"),
    "CHK-166": ("Employee Annual Leave Setting Incorrect", "Low"),
    "CHK-167": ("Employee Long Service Leave Entitlements Missing", "Medium"),
    "CHK-168": ("Check Earnings Rate Category", "Low"),
}
# Catalogue checks a built one covers: CHK-149 covers CHK-160 (monthly leave balance warning), CHK-157 covers CHK-158
# (super liability payments), CHK-164 covers CHK-165 (monthly super contribution check).
PAY_ALSO = {"CHK-149": ("CHK-160",), "CHK-157": ("CHK-158",), "CHK-164": ("CHK-165",)}
# Readings with thresholds of our own, until the content owner confirms them.
PAY_HEURISTIC = ("CHK-156", "CHK-164", "CHK-166", "CHK-168")
PAY_SEVERITY_RANK = {"High": 0, "Medium": 1, "Low": 2}
# The national minimum wage an hour, from the first pay period starting on or after the date (Fair Work Commission annual
# wage reviews; $26.44 from 1 July 2026). Add each year's on 1 July. A casual's floor adds the 25% casual loading.
PAY_WAGE_FLOOR = (("2025-07-01", 24.95), ("2026-07-01", 26.44))
PAY_CASUAL_LOADING = 0.25
PAY_ADULT_AGE = 21
# The super guarantee rate (% of ordinary time earnings).
PAY_SG_RATE = (("2025-07-01", 12.0),)
# Payday Super (from 1 July 2026): super has to reach the fund within 7 business days of pay day. Before it, 28 days after
# the end of the quarter.
PAY_PAYDAY_SUPER = "2026-07-01"
PAY_SUPER_BUSINESS_DAYS = 7
# Business days skip weekends and the national public holidays; one more day allows for a state holiday.
PAY_SUPER_SLACK_DAYS = 1
PAY_TFN_DAYS = 28
PAY_FULL_TIME_HOURS = 38.0
PAY_TERMINATED_DAYS = 400
PAY_OTE_TYPES = ("ORDINARYTIMEEARNINGS", "ALLOWANCE", "BONUSESANDCOMMISSIONS")
# CHK-168 (heuristic): a pay item's name against its STP earnings type.
PAY_ITEM_TYPES = ((r"(?i)overtime|\bo/?t\b|time and a half|double time", "OVERTIMEEARNINGS", "overtime"),
                  (r"(?i)allowance", "ALLOWANCE", "an allowance"),
                  (r"(?i)bonus|commission", "BONUSESANDCOMMISSIONS", "a bonus or commission"),
                  (r"(?i)director", "DIRECTORSFEES", "directors' fees"),
                  (r"(?i)parental", "PAIDPARENTALLEAVE", "paid parental leave"),
                  (r"(?i)workers'? ?comp", "WORKERSCOMPENSATION", "workers' compensation"))


def pay_from(table, day):
    """The value in effect on a day, from a table of (from date, value)."""
    out = None
    for start, value in table:
        if str(day) >= start:
            out = value
    return out


def pay_money(amount):
    return "$" + "{:,.2f}".format(amount or 0.0)


def pay_easter(year):
    """Easter Sunday (the Gregorian computus)."""
    a, b, c = year % 19, year // 100, year % 100
    d, e = b // 4, b % 4
    g = (8 * b + 13) // 25
    h = (19 * a + b - d - g + 15) % 30
    i, k = c // 4, c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 19 * l) // 433
    month = (h + l - 7 * m + 90) // 25
    return datetime.date(year, month, (h + l - 7 * m + 33 * month + 19) % 32)


def pay_national_holidays(year):
    """Australia's national public holidays in a year, with the Monday (or Tuesday) a weekend one moves to."""
    days = set()
    for month, day in ((1, 1), (1, 26)):
        d = datetime.date(year, month, day)
        days.add(d if d.weekday() < 5 else d + datetime.timedelta(days=7 - d.weekday()))
    christmas, boxing = datetime.date(year, 12, 25), datetime.date(year, 12, 26)
    days.add(christmas if christmas.weekday() < 5 else christmas + datetime.timedelta(days=7 - christmas.weekday()))
    moved = boxing + datetime.timedelta(days=7 - boxing.weekday()) if boxing.weekday() >= 5 else boxing
    days.add(moved + datetime.timedelta(days=1) if moved in days else moved)
    easter = pay_easter(year)
    days |= {easter - datetime.timedelta(days=2), easter + datetime.timedelta(days=1), datetime.date(year, 4, 25)}
    return days


def pay_day_text(value):
    d = pay_date(value)
    return d.strftime("%d %b %Y").lstrip("0") if d else "no date"


def pay_finding(check, name, text, amount=None):
    title, severity = PAY_CHECKS[check]
    return {"check": check, "title": title, "severity": severity, "name": name, "text": text,
            "amount": round(amount, 2) if amount is not None else None, "heuristic": check in PAY_HEURISTIC,
            "also": list(PAY_ALSO.get(check, ()))}


def pay_leave_kind(leave_type):
    """annual, personal, long service or other: the leave type's STP category, or its name when it has none."""
    code = str((leave_type or {}).get("LeaveCategoryCode") or "").upper()
    name = str((leave_type or {}).get("Name") or "").lower()
    if code == "ANNUALLEAVE" or (not code and "annual" in name):
        return "annual"
    if code == "LONGSERVICELEAVE" or (not code and "long service" in name):
        return "long service"
    if code == "PERSONALSICKCARERSLEAVE" or (not code and re.search(r"personal|sick|carer", name)):
        return "personal"
    return "other"


def pay_line_amount(line):
    amount = bk_money(line.get("Amount"))
    if amount is None:
        amount = bk_money(line.get("FixedAmount"))
    if amount is None:
        amount = (bk_money(line.get("RatePerUnit")) or 0.0) * (bk_money(line.get("NumberOfUnits")) or 0.0)
    return amount or 0.0


def pay_hourly(e, earnings):
    """(ordinary hourly rate, hours a week or None) from the employee's pay template line for their ordinary earnings
    rate (a salary: annual / 52 / hours a week), else the earnings rate's own; None when it isn't paid by the hour."""
    oid = e.get("OrdinaryEarningsRateID")
    er = earnings.get(oid) or {}
    if str(er.get("TypeOfUnits") or "Hours").strip().lower() not in ("hours", "hour", "hrs"):
        return None
    for line in (e.get("PayTemplate") or {}).get("EarningsLines") or []:
        if line.get("EarningsRateID") == oid:
            salary, weekly = bk_money(line.get("AnnualSalary")), bk_money(line.get("NumberOfUnitsPerWeek"))
            if salary and weekly:
                return salary / 52 / weekly, weekly
            rate = bk_money(line.get("RatePerUnit"))
            if rate:
                return rate, weekly
    rate = bk_money(er.get("RatePerUnit"))
    return (rate, None) if rate else None


def pay_age(e, day):
    dob = pay_date(e.get("DateOfBirth"))
    return (day.year - dob.year - ((day.month, day.day) < (dob.month, dob.day))) if dob else None


def pay_business_days_after(day, n):
    d = day
    while n > 0:
        d += datetime.timedelta(days=1)
        if d.weekday() < 5 and d not in pay_national_holidays(d.year):
            n -= 1
    return d


def pay_super_due(paid_on):
    """When super for a pay day has to reach the fund."""
    if str(paid_on) >= PAY_PAYDAY_SUPER:
        return pay_business_days_after(paid_on, PAY_SUPER_BUSINESS_DAYS + PAY_SUPER_SLACK_DAYS)
    quarter_end = ((paid_on.month - 1) // 3 + 1) * 3
    return datetime.date(paid_on.year + (1 if quarter_end == 12 else 0), quarter_end % 12 + 1, 28)


def pay_super_overdue(owed, posted, today):
    """(overdue, not yet due) on the super liability account: what's owed, less the super of posted pay runs whose super
    isn't due yet."""
    pending = 0.0
    for r in posted:
        paid_on = pay_date(r.get("PaymentDate"))
        if paid_on and pay_super_due(paid_on) >= today:
            pending += bk_money(r.get("Super")) or 0.0
    return round(owed - pending, 2), round(pending, 2)


def pay_ote(slip, earnings):
    """A payslip's ordinary time earnings: ordinary hours, allowances and bonuses, and leave taken, unless the pay item
    is exempt from super."""
    total = 0.0
    for kind in ("EarningsLines", "TimesheetEarningsLines", "LeaveEarningsLines"):
        for ln in slip.get(kind) or []:
            er = earnings.get(ln.get("EarningsRateID")) or {}
            if not er.get("IsExemptFromSuper") and (kind == "LeaveEarningsLines" or str(er.get("EarningsType") or "") in PAY_OTE_TYPES):
                total += pay_line_amount(ln)
    return round(total, 2)


def pay_catalogue_checks(staff, terminated, earnings, leave_types, draft, payslips, start, end, today):
    """The catalogue's payroll checks for one pay run. staff: the active employees on the pay calendar (full records);
    terminated: the terminated ones; earnings / leave_types: the pay items by id; draft: the draft pay run or None;
    payslips: the draft's payslips in full, by employee id. Returns findings, most severe first."""
    out = []
    everyone = {str(e.get("EmployeeID")): e for e in list(staff) + list(terminated) if e.get("EmployeeID")}
    paid = {str(s["EmployeeID"]): s for s in (draft or {}).get("Payslips") or [] if s.get("EmployeeID")}

    # CHK-147: terminated in the last year or so without the cessation reason Single Touch Payroll reports.
    for e in terminated:
        gone = pay_date(e.get("TerminationDate"))
        if gone and (today - gone).days <= PAY_TERMINATED_DAYS and not str(e.get("TerminationReason") or "").strip():
            out.append(pay_finding("CHK-147", pay_name(e), "was terminated on " + pay_day_text(e.get("TerminationDate"))
                                   + " with no cessation reason. Single Touch Payroll needs one (voluntary, ill health, deceased, "
                                   "redundancy, dismissal, contract cessation or transfer): add it in Xero Payroll and send an "
                                   "update event."))

    # On this pay run: terminated before the period (CHK-155), starting after it (CHK-153), reimbursed more than paid (CHK-151).
    for eid, s in paid.items():
        e = everyone.get(eid) or {}
        name = pay_name(e) or pay_name(s)
        wages = bk_money(s.get("Wages")) or 0.0
        gone = pay_date(e.get("TerminationDate"))
        if gone and gone < start and wages > 0:
            out.append(pay_finding("CHK-155", name, "was terminated on " + pay_day_text(e.get("TerminationDate")) + " but is paid "
                                   + pay_money(wages) + " in wages on this pay run. Check it's their final pay, or take them off it.",
                                   amount=wages))
        began = pay_date(e.get("StartDate"))
        if began and began > end and wages > 0:
            out.append(pay_finding("CHK-153", name, "is paid " + pay_money(wages) + " on this pay run but starts on "
                                   + pay_day_text(e.get("StartDate")) + ", after the period ends. Check the start date, or take "
                                   "them off it.", amount=wages))
        refund = bk_money(s.get("Reimbursements")) or 0.0
        if refund > 0 and refund > wages:
            out.append(pay_finding("CHK-151", name, "is reimbursed " + pay_money(refund) + " on this pay run, more than their wages ("
                                   + pay_money(wages) + "). Check the reimbursement lines and the receipts behind them.", amount=refund))

    floor = pay_from(PAY_WAGE_FLOOR, start)
    sg_rate = pay_from(PAY_SG_RATE, start)
    banks = {}
    for e in staff:
        name = pay_name(e)
        eid = str(e.get("EmployeeID") or "")
        basis = pay_basis(e)
        age = pay_age(e, end)
        tax = e.get("TaxDeclaration") or {}
        kinds = {}
        for ln in (e.get("PayTemplate") or {}).get("LeaveLines") or []:
            if isinstance(ln, dict):
                kinds.setdefault(pay_leave_kind(leave_types.get(ln.get("LeaveTypeID"))), []).append(ln)
        slip = payslips.get(eid) or {}
        hourly = pay_hourly(e, earnings)

        # CHK-148: "TFN applied for" more than 28 days after starting. The number itself is never read out.
        began = pay_date(e.get("StartDate"))
        if str(tax.get("TFNExemptionType") or "").upper() == "PENDING" and began and (today - began).days > PAY_TFN_DAYS:
            out.append(pay_finding("CHK-148", name, "still has no tax file number " + str((today - began).days) + " days after "
                                   "starting (it was applied for). After 28 days, withhold tax at the no-TFN rate until they "
                                   "provide it."))
        # CHK-149 (covers CHK-160): a negative leave balance.
        for b in e.get("LeaveBalances") or []:
            units = bk_money(b.get("NumberOfUnits"))
            if units is not None and units < -0.005:
                out.append(pay_finding("CHK-149", name, "has a negative " + str(b.get("LeaveName") or "leave") + " balance ("
                                       + str(round(units, 2)) + " " + str(b.get("TypeOfUnits") or "hours").lower()
                                       + "). Check the leave taken and the accrual."))
        # CHK-150: a casual accruing annual leave (casuals are paid a loading instead).
        accruing = [ln for ln in kinds.get("annual") or [] if str(ln.get("CalculationType") or "") != "NOCALCULATIONREQUIRED"]
        accrued = [a for a in slip.get("LeaveAccrualLines") or []
                   if pay_leave_kind(leave_types.get(a.get("LeaveTypeID"))) == "annual" and (bk_money(a.get("NumberOfUnits")) or 0) > 0]
        if basis == "casual" and (accruing or accrued):
            out.append(pay_finding("CHK-150", name, "is a casual but accrues annual leave. Casuals are paid a loading instead: "
                                   "take the annual leave line off their pay template."))
        if basis in ("full-time", "part-time"):
            # CHK-161: no annual or personal/carer's leave accruing; CHK-167: no long service leave.
            missing = [words for kind, words in (("annual", "annual leave"), ("personal", "personal/carer's leave")) if not kinds.get(kind)]
            if missing:
                out.append(pay_finding("CHK-161", name, "is " + basis + " but has no " + " or ".join(missing) + " in their pay "
                                       "template, so it isn't accruing. Add the leave line(s) in Xero Payroll."))
            if not kinds.get("long service"):
                out.append(pay_finding("CHK-167", name, "is " + basis + " but has no long service leave in their pay template, so "
                                       "it isn't accruing. Add it if the state's long service leave law applies to them."))
            # CHK-166 (heuristic): annual leave set not to accrue, or a full-timer's that isn't 4 (or 5) weeks a year.
            weekly = (hourly[1] if hourly and hourly[1] else None) or PAY_FULL_TIME_HOURS
            for ln in kinds.get("annual") or []:
                calc = str(ln.get("CalculationType") or "")
                units = bk_money(ln.get("AnnualNumberOfUnits"))
                if calc == "NOCALCULATIONREQUIRED":
                    out.append(pay_finding("CHK-166", name, "is " + basis + " but their annual leave is set to no calculation, "
                                           "so none accrues. Set it to accrue (4 weeks a year, pro rata for part-time)."))
                elif basis == "full-time" and units and calc in ("FIXEDAMOUNTEACHPERIOD", "ENTERRATEINPAYTEMPLATE") \
                        and min(abs(units - 4 * weekly), abs(units - 5 * weekly)) > 0.5:
                    out.append(pay_finding("CHK-166", name, "accrues " + str(round(units, 2)) + " hours of annual leave a year; 4 "
                                           "weeks at " + str(round(weekly, 2)) + " hours a week is " + str(round(4 * weekly, 2))
                                           + " (5 weeks, " + str(round(5 * weekly, 2)) + ", for a shift worker). Check the leave line."))
        # CHK-156 (heuristic): an adult's ordinary rate below the national minimum wage.
        if floor and hourly and basis in ("full-time", "part-time", "casual") and age is not None and age >= PAY_ADULT_AGE \
                and str(e.get("EmploymentType") or "EMPLOYEE").upper() != "CONTRACTOR":
            least = round(floor * (1 + PAY_CASUAL_LOADING), 2) if basis == "casual" else floor
            if hourly[0] < least - 0.005:
                hours = sum(bk_money(ln.get("NumberOfUnits")) or 0.0
                            for ln in (slip.get("EarningsLines") or []) + (slip.get("TimesheetEarningsLines") or [])
                            if ln.get("EarningsRateID") == e.get("OrdinaryEarningsRateID"))
                out.append(pay_finding("CHK-156", name, "is paid " + pay_money(hourly[0]) + " an hour, below the national minimum "
                                       "wage of " + pay_money(least) + (" for a casual (with the 25% loading)" if basis == "casual" else "")
                                       + ". Check their award or agreement rate (apprentices, trainees and supported wages differ).",
                                       amount=(least - hourly[0]) * hours if hours else None))
        # CHK-164 (heuristic, covers CHK-165): the super guarantee on this pay run against the ordinary time earnings.
        if slip and sg_rate and not (age is not None and age < 18):
            ote = pay_ote(slip, earnings)
            sgc = sum(bk_money(x.get("Amount")) or 0.0 for x in slip.get("SuperannuationLines") or []
                      if str(x.get("ContributionType") or "") == "SGC")
            want = round(ote * sg_rate / 100, 2)
            if ote > 0 and want - sgc > max(1.0, want * 0.02):
                out.append(pay_finding("CHK-164", name, "has " + pay_money(sgc) + " super guarantee on this pay run, but "
                                       + ("%g" % sg_rate) + "% of their ordinary time earnings (" + pay_money(ote) + ") is "
                                       + pay_money(want) + ". Check their super line (someone over the maximum contribution base "
                                       "can have less).", amount=want - sgc))
        # CHK-159: being paid with details its reporting needs missing. Only the gap is named, never the detail.
        if eid in paid or not draft:
            gaps = []
            if not tax or not (str(tax.get("TaxFileNumber") or "").strip() or tax.get("TFNExemptionType")):
                gaps.append("tax file number declaration")
            if not e.get("SuperMemberships") and not (age is not None and age < 18):
                gaps.append("super fund")
            if not e.get("BankAccounts"):
                gaps.append("bank account")
            if not str((e.get("HomeAddress") or {}).get("AddressLine1") or "").strip():
                gaps.append("home address")
            if e.get("IsSTP2Qualified") is False:
                gaps.append("Single Touch Payroll Phase 2 update")
            if gaps:
                out.append(pay_finding("CHK-159", name, ("is on this pay run" if draft else "is on this pay calendar")
                                       + " but their Xero Payroll record has no " + ", no ".join(gaps)
                                       + ". Add them before the pay run is posted."))
        # CHK-162: one bank account for two employees, compared in memory as a hash only.
        for a in e.get("BankAccounts") or []:
            digits = re.sub(r"\D", "", str(a.get("BSB") or "")) + "|" + re.sub(r"\D", "", str(a.get("AccountNumber") or ""))
            if len(digits) > 4:
                banks.setdefault(hashlib.sha256(digits.encode("utf-8")).hexdigest(), set()).add(name)
    for names in banks.values():
        if len(names) > 1:
            who = sorted(names)
            out.append(pay_finding("CHK-162", ", ".join(who), " and ".join(who) + " are paid into the same bank account. Check "
                                   "with them it's right: a shared account is rare, and can be a sign of fraud."))

    # CHK-168 (heuristic): pay items in use whose name and STP earnings type disagree, or an allowance with no type.
    used = set()
    for s in payslips.values():
        used |= {ln.get("EarningsRateID") for ln in (s.get("EarningsLines") or []) + (s.get("TimesheetEarningsLines") or [])}
    for e in staff:
        used |= {ln.get("EarningsRateID") for ln in (e.get("PayTemplate") or {}).get("EarningsLines") or []}
    for rid in sorted(r for r in used if r in earnings):
        er = earnings[rid]
        item, kind = str(er.get("Name") or ""), str(er.get("EarningsType") or "")
        clash = next(((want, words) for pattern, want, words in PAY_ITEM_TYPES if re.search(pattern, item)), None)
        if clash and kind != clash[0]:
            out.append(pay_finding("CHK-168", "", "Pay item " + item + " sounds like " + clash[1] + " but is set up as "
                                   + (kind.lower().replace("earnings", " earnings") if kind else "no earnings type") + ". Check its "
                                   "earnings type: Single Touch Payroll reports it under that."))
        elif kind == "ALLOWANCE" and not er.get("AllowanceType"):
            out.append(pay_finding("CHK-168", "", "Pay item " + item + " is an allowance with no allowance type, which Single "
                                   "Touch Payroll needs. Set its type in Xero Payroll."))
    return sorted(out, key=lambda f: (PAY_SEVERITY_RANK[f["severity"]], f["check"], f["name"]))
