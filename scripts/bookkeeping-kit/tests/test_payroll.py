"""Payroll phase 1: open the pay run, check it (Deputy vs Xero Payroll, rates, penalties, changes, reimbursements),
build the review pack, internal review, client approval. Nothing is ever written to Xero Payroll."""
import datetime
import os
import io
import json
import unittest
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook

from harness import run_flow, xero_ms_date
from test_flows import AUTOS, BOOKKEEPER, T_ACME, Base, source_id

try:
    SYD = ZoneInfo("Australia/Sydney")
except Exception:  # no tz database here: the kit falls back to UTC+10 the same way (bk_tz)
    SYD = datetime.timezone(datetime.timedelta(hours=10))


def at(day, hh, mm):
    """A run time on the client's clock, as the scheduler would pass it."""
    return datetime.datetime(day.year, day.month, day.day, hh, mm, tzinfo=SYD).isoformat()


PREP = datetime.date(2026, 10, 5)
START, END, PAY = datetime.date(2026, 9, 21), datetime.date(2026, 10, 4), datetime.date(2026, 10, 6)
CHECK = "Bookkeeping: Check pay run (Xero Payroll + Deputy)"
CHECK_XERO = "Bookkeeping: Check pay run (Xero Payroll)"
PACK = "Bookkeeping: Build payroll review pack (Xero Payroll)"
REVIEW = "Bookkeeping: Apply payroll review"
STARTRUNS = "Bookkeeping: Start pay runs (Xero Payroll)"
OFFICER = {"id": "user-officer", "label": "Pat Officer"}
SAM = {"id": "user-sam", "label": "Sam Reviewer"}


def epoch(day, hh, mm):
    return int(datetime.datetime(day.year, day.month, day.day, hh, mm, tzinfo=SYD).timestamp())


def workbook_bytes(sheets):
    book = Workbook()
    blank = book.active
    for title, rows in sheets.items():
        ws = book.create_sheet(title=title)
        for r in rows:
            ws.append(r)
    book.remove(blank)
    buf = io.BytesIO()
    book.save(buf)
    return buf.getvalue()


PAY_GUIDE = {
    "rates": [["award", "classification", "basis", "age_from", "age_to", "band", "rate", "effective_from"],
              ["GRIA", "Level 1", "casual", 18, None, "ordinary", 30.00, "2026-07-01"],
              ["GRIA", "Level 1", "casual", 18, None, "saturday", 36.00, "2026-07-01"],
              ["GRIA", "Level 1", "casual", 18, None, "early", 33.00, "2026-07-01"],
              ["GRIA", "Level 1", "casual", 16, 17, "ordinary", 24.00, "2026-07-01"]],
    "employees": [["employee", "award", "classification", "state", "special_rate", "special_until", "notes"],
                  ["Casey Casual", "GRIA", "Level 1", "NSW", None, None, ""],
                  ["Sam Special", "GRIA", "Level 1", "NSW", 35.00, "2026-09-28", "Flat rate until 28 Sep"]],
    "rules": [["award", "day_start", "day_end", "ot_after_hours", "ot_first_hours"], ["GRIA", "07:00", "18:00", 9, 3]],
    "holidays": [["date", "state", "name"]],
    "pay_items": [["pay_item", "band"], ["Casual Ordinary Hours", "ordinary"], ["Casual Saturday", "saturday"],
                  ["Casual before 7am", "early"]],
}


class PayrollTests(Base):
    def setUp(self):
        super().setUp()
        self.setup_practice()
        run_flow(self.w, AUTOS["Bookkeeping: Start setup"], {"kind": "manual", "actor_user_id": BOOKKEEPER},
                 {"payroll_names": "Acme Plumbing"})
        item = next(i for i in self.w.items.values() if i["title"] == "Payroll settings for Acme Plumbing")
        self.settings_item = item
        self.submit("Bookkeeping: Apply setup (Xero)", "Client payroll settings", item["id"], {
            "q_client_name": "Acme Plumbing", "q_calendar": "Fortnightly", "q_prep_day": "Monday", "q_prep_time": "09:00",
            "q_states": "NSW", "q_timesheets": "Deputy (the client exports them to Xero)", "q_ts_owner": "Robin",
            "q_ts_email": "robin@acme.example", "q_chase_after": "60", "q_officer": OFFICER, "q_reviewer": SAM,
            "q_approvers": "Robin and Jordan (owners@acme.example)", "q_who_posts": "The client"})
        self.seed_payroll()

    def seed_payroll(self, with_timesheets=True):
        o = self.w.payroll.org(T_ACME)
        o["calendars"] = [{"PayrollCalendarID": "CAL1", "Name": "Fortnightly", "CalendarType": "FORTNIGHTLY",
                           "StartDate": xero_ms_date(START), "PaymentDate": xero_ms_date(PAY)}]
        o["earnings"] = [{"EarningsRateID": "ORD", "Name": "Casual Ordinary Hours", "RatePerUnit": 30.0},
                         {"EarningsRateID": "SAT", "Name": "Casual Saturday", "RatePerUnit": 36.0},
                         {"EarningsRateID": "EARLY", "Name": "Casual before 7am", "RatePerUnit": 33.0}]

        def emp(eid, first, last, basis, dob, rate=None):
            e = {"EmployeeID": eid, "FirstName": first, "LastName": last, "Status": "ACTIVE", "PayrollCalendarID": "CAL1",
                 "DateOfBirth": xero_ms_date(dob), "TaxDeclaration": {"EmploymentBasis": basis},
                 "OrdinaryEarningsRateID": "ORD", "HomeAddress": {"Region": "NSW"}}
            if rate is not None:
                e["PayTemplate"] = {"EarningsLines": [{"EarningsRateID": "ORD", "RatePerUnit": rate}]}
            return e
        o["employees"] = [emp("E1", "Casey", "Casual", "CASUAL", datetime.date(2008, 9, 25), 30.0),
                          emp("E2", "Fran", "Fulltime", "FULLTIME", datetime.date(1985, 1, 1)),
                          emp("E3", "Sam", "Special", "CASUAL", datetime.date(1990, 5, 5), 35.0)]

        def sheet(eid, units):
            return {"TimesheetID": "T" + eid, "EmployeeID": eid, "StartDate": xero_ms_date(START), "EndDate": xero_ms_date(END),
                    "Status": "DRAFT", "TimesheetLines": [{"EarningsRateID": "ORD", "NumberOfUnits": units}]}
        days = [0.0] * 14
        casey = list(days); casey[0] = 8.0; casey[5] = 5.0            # Mon 21 and Sat 26, both as ordinary hours
        fran = [7.6 if i % 7 < 5 else 0.0 for i in range(14)]        # salaried: shouldn't have a timesheet
        sam = list(days); sam[1] = 6.0                               # Tue 22
        o["timesheets"] = [sheet("E1", casey), sheet("E2", fran), sheet("E3", sam)] if with_timesheets else []
        o["pay_runs"] = [{"PayRunID": "PR1", "PayrollCalendarID": "CAL1", "PayRunStatus": "DRAFT",
                          "PayRunPeriodStartDate": xero_ms_date(START), "PayRunPeriodEndDate": xero_ms_date(END),
                          "Wages": 2950.0, "Tax": 410.0, "Super": 354.0, "Deductions": 0, "Reimbursement": 20.0, "NetPay": 2560.0,
                          "Payslips": [{"EmployeeID": "E1", "FirstName": "Casey", "LastName": "Casual", "Wages": 390, "Tax": 20,
                                        "Super": 46.8, "Reimbursements": 0, "NetPay": 370},
                                       {"EmployeeID": "E2", "FirstName": "Fran", "LastName": "Fulltime", "Wages": 2350, "Tax": 380,
                                        "Super": 282, "Reimbursements": 0, "NetPay": 1970},
                                       {"EmployeeID": "E3", "FirstName": "Sam", "LastName": "Special", "Wages": 210, "Tax": 10,
                                        "Super": 25.2, "Reimbursements": 20, "NetPay": 220}]}]
        d = self.w.deputy
        d.employees = [{"Id": 1, "DisplayName": "Casey Casual"}, {"Id": 3, "DisplayName": "Sam Special"},
                       {"Id": 9, "DisplayName": "New Person"}]

        def shift(emp_id, day, sh, sm, eh, em, total):
            return {"Id": len(d.timesheets) + 1, "Employee": emp_id, "Date": str(day) + "T00:00:00+10:00",
                    "StartTime": epoch(day, sh, sm), "EndTime": epoch(day, eh, em), "TotalTime": total}
        d.timesheets = [shift(1, datetime.date(2026, 9, 21), 6, 0, 14, 30, 8.0),
                        shift(1, datetime.date(2026, 9, 23), 9, 0, 13, 0, 4.0),
                        shift(1, datetime.date(2026, 9, 26), 9, 0, 14, 0, 5.0),
                        shift(3, datetime.date(2026, 9, 22), 9, 0, 15, 0, 6.0),
                        shift(9, datetime.date(2026, 9, 24), 10, 0, 14, 0, 4.0)]
        self.w.files["Bookkeeping Payroll/acme-plumbing/pay-guide.xlsx"] = workbook_bytes(PAY_GUIDE)
        self.w.files["Bookkeeping Payroll/acme-plumbing/Pop expense September.xlsx"] = workbook_bytes(
            {"Sheet1": [["Pop expense claims", None, None], ["Employee", "Type", "Amount"], ["Casey Casual", "Mileage", 45.50],
                        ["Sam Special", "Other", 20.00]]})

    # ── helpers ──
    def start_runs(self, now, client=None):
        return run_flow(self.w, AUTOS[STARTRUNS], {"kind": "schedule", "actor_user_id": BOOKKEEPER, "now": now},
                        {"client_name": client} if client else None)

    def open_run(self):
        self.start_runs(at(PREP, 9, 30))
        return next(i for i in self.w.items.values() if i["title"].startswith("Pay run: Acme Plumbing"))

    def run_record(self):
        return json.loads(self.w.files["Bookkeeping/PayRuns/acme-plumbing-2026-09-21.json"])

    def check(self, item, auto=CHECK):
        return self.submit(auto, "Pay run", item["id"], {"q_action": "Run the checks"}, actor=OFFICER["id"])

    # ── settings ──
    def test_payroll_settings_are_saved_with_a_pay_guide_template(self):
        ps = json.loads(self.w.files["Bookkeeping/Payroll/acme-plumbing.json"])
        self.assertEqual((ps["calendar"], ps["prep_day"], ps["timesheets"], ps["who_posts"]), ("Fortnightly", "Monday", "deputy", "client"))
        self.assertEqual((ps["officer_id"], ps["reviewer_id"], ps["states"]), ("user-officer", "user-sam", ["NSW"]))
        self.assertEqual(self.w.items[self.settings_item["id"]]["status"], "done")
        self.assertIn("pay-guide.xlsx", self.w.items[self.settings_item["id"]]["comments"][-1])

    def test_a_new_client_gets_a_pay_guide_template_but_a_filled_one_is_never_overwritten(self):
        filled = self.w.files["Bookkeeping Payroll/acme-plumbing/pay-guide.xlsx"]
        self.submit("Bookkeeping: Apply setup (Xero)", "Client payroll settings", self.settings_item["id"], {
            "q_client_name": "Acme Plumbing", "q_calendar": "Fortnightly", "q_prep_day": "Monday", "q_timesheets": "Entered in Xero Payroll",
            "q_officer": OFFICER, "q_reviewer": SAM, "q_who_posts": "Our practice"})
        self.assertEqual(self.w.files["Bookkeeping Payroll/acme-plumbing/pay-guide.xlsx"], filled)
        del self.w.files["Bookkeeping Payroll/acme-plumbing/pay-guide.xlsx"]
        self.submit("Bookkeeping: Apply setup (Xero)", "Client payroll settings", self.settings_item["id"], {
            "q_client_name": "Acme Plumbing", "q_calendar": "Fortnightly", "q_prep_day": "Monday", "q_timesheets": "Entered in Xero Payroll",
            "q_officer": OFFICER, "q_reviewer": SAM, "q_who_posts": "Our practice"})
        book = load_workbook(io.BytesIO(self.w.files["Bookkeeping Payroll/acme-plumbing/pay-guide.xlsx"]))
        self.assertEqual(book.sheetnames, ["rates", "employees", "rules", "holidays", "pay_items"])

    def test_payroll_settings_need_the_bookkeeping_client_first(self):
        steps = self.submit("Bookkeeping: Apply setup (Xero)", "Client payroll settings", self.settings_item["id"], {
            "q_client_name": "Nobody Pty Ltd", "q_calendar": "Weekly", "q_prep_day": "Friday", "q_timesheets": "Entered in Xero Payroll",
            "q_officer": OFFICER, "q_reviewer": SAM, "q_who_posts": "The client"})
        self.assertEqual(steps["apply_setup"]["outputs"]["outcome"], "payroll_client_missing")

    # ── starting the pay run ──
    def test_the_pay_run_opens_on_the_preparation_day_once(self):
        self.start_runs(at(PREP, 8, 30))
        self.assertFalse(any(i["title"].startswith("Pay run:") for i in self.w.items.values()))
        self.start_runs(at(PREP, 9, 30))
        self.start_runs(at(PREP, 10, 30))
        runs = [i for i in self.w.items.values() if i["title"].startswith("Pay run:")]
        self.assertEqual(len(runs), 1)
        item = runs[0]
        self.assertEqual(item["title"], "Pay run: Acme Plumbing, 2026-09-21 to 2026-10-04, pay day 2026-10-06")
        self.assertEqual((item["assigneeIds"], item["dueDate"]), (["user-officer"], "2026-10-06"))
        self.assertEqual(item["formId"], source_id(STARTRUNS, "Pay run"))
        self.assertIn("Timesheets are in Xero Payroll (3)", item["comments"][0])

    def test_late_deputy_timesheets_are_chased_once(self):
        self.seed_payroll(with_timesheets=False)
        self.start_runs(at(PREP, 9, 30))
        item = next(i for i in self.w.items.values() if i["title"].startswith("Pay run:"))
        self.assertEqual(item["comments"], [])
        self.start_runs(at(PREP, 10, 5))
        self.start_runs(at(PREP, 11, 5))
        chases = [c for c in self.w.items[item["id"]]["comments"] if "Draft follow-up to Robin (robin@acme.example)" in c]
        self.assertEqual(len(chases), 1)

    def test_a_missing_pay_calendar_is_reported_not_opened(self):
        self.w.payroll.org(T_ACME)["calendars"][0]["Name"] = "Weekly"
        steps = self.start_runs(at(PREP, 9, 30))
        self.assertIn('no pay calendar called "Fortnightly"', steps["start_runs"]["outputs"]["summary"])

    # ── the check ──
    def test_the_check_finds_what_the_payroll_officer_looks_for(self):
        item = self.open_run()
        steps = self.check(item)
        f = steps["check"]["outputs"]["findings"]
        self.assertEqual(steps["check"]["outputs"]["problem"], "")
        self.assertEqual(f["missing_starters"], ["New Person"])
        self.assertEqual([r["name"] for r in f["salaried_timesheets"]], ["Fran Fulltime"])
        self.assertEqual(f["hours"], [{"name": "Casey Casual", "date": "2026-09-23", "deputy": 4.0, "xero": 0}])
        bands = {r["band"]: r for r in f["bands"] if r["name"] == "Casey Casual"}
        self.assertEqual(set(bands), {"ordinary", "early", "saturday"})
        self.assertAlmostEqual(bands["saturday"]["expected"], 5.0)
        self.assertAlmostEqual(bands["early"]["expected"], round(8 / 8.5, 2))
        self.assertEqual(f["rates"], [])  # Sam's flat rate is in effect until 28 Sep, inside this period
        changes = {r["name"]: r["text"] for r in f["changes"]}
        # CHK-163, without the birthday itself (it's the date of birth, which payroll findings never show).
        self.assertIn("turns 18 during this pay period", changes["Casey Casual"])
        self.assertNotIn("2026-09-25", json.dumps(f))
        self.assertIn("special arrangement ends 2026-09-28", changes["Sam Special"])
        self.assertEqual(f["reimbursements"], [{"name": "Casey Casual", "file": 45.5, "pay_run": 0}])
        self.assertEqual(f["totals"]["NetPay"], 2560.0)

    def test_an_arrangement_that_ended_before_the_period_is_a_rate_error(self):
        guide = json.loads(json.dumps(PAY_GUIDE))
        guide["employees"][2][5] = "2026-08-17"
        self.w.files["Bookkeeping Payroll/acme-plumbing/pay-guide.xlsx"] = workbook_bytes(guide)
        f = self.check(self.open_run())["check"]["outputs"]["findings"]
        self.assertEqual([(r["name"], r["band"], r["expected"], r["xero"]) for r in f["rates"]], [("Sam Special", "ordinary", 30.0, 35.0)])
        self.assertIn("ended 2026-08-17 but Xero still pays the special rate", {r["name"]: r["text"] for r in f["changes"]}["Sam Special"])

    def test_a_casual_missing_from_the_pay_guide_is_listed(self):
        guide = json.loads(json.dumps(PAY_GUIDE))
        del guide["employees"][1]
        self.w.files["Bookkeeping Payroll/acme-plumbing/pay-guide.xlsx"] = workbook_bytes(guide)
        f = self.check(self.open_run())["check"]["outputs"]["findings"]
        self.assertEqual(f["unclassified"], ["Casey Casual"])

    def test_the_check_reports_and_drafts_one_query_email(self):
        item = self.open_run()
        steps = self.check(item)
        self.assertGreater(steps["report"]["outputs"]["open_items"], 0)
        comment = self.w.items[item["id"]]["comments"][-1]
        self.assertIn("Missing in Xero Payroll (hours in Deputy): **1**", comment)
        self.assertIn("Salaried staff with timesheets: **1**", comment)
        self.assertIn("Draft query email to Robin (robin@acme.example)", comment)
        self.assertIn("[Open the full check]", comment)
        report = next(a for a in self.w.artifacts if a["title"].startswith("Payroll check: Acme Plumbing"))
        self.assertIn("New Person", report["content"])
        self.assertEqual(self.run_record()["status"], "checked")
        self.assertEqual(sum(1 for c in self.w.ai_calls if c[0] == "generate"), 1)

    def test_the_xero_only_check_never_calls_deputy_and_says_so(self):
        item = self.open_run()
        steps = self.check(item, auto=CHECK_XERO)
        self.assertEqual(self.w.deputy.calls, [])
        self.assertIn("Check pay run (Xero Payroll + Deputy)", " ".join(steps["check"]["outputs"]["findings"]["notes"]))
        self.assertEqual([r["name"] for r in steps["check"]["outputs"]["findings"]["salaried_timesheets"]], ["Fran Fulltime"])

    def test_without_a_pay_guide_rates_are_skipped_with_a_note(self):
        del self.w.files["Bookkeeping Payroll/acme-plumbing/pay-guide.xlsx"]
        item = self.open_run()
        f = self.check(item)["check"]["outputs"]["findings"]
        self.assertEqual((f["rates"], f["bands"]), ([], []))
        self.assertIn("There's no pay guide", " ".join(f["notes"]))

    def test_a_manual_run_checks_the_clients_open_pay_run(self):
        self.open_run()
        steps = run_flow(self.w, AUTOS[CHECK], {"kind": "manual", "actor_user_id": BOOKKEEPER}, {"client_name": "Acme Plumbing"})
        self.assertTrue(steps["load_run"]["outputs"]["should_check"])
        self.assertEqual(steps["check"]["outputs"]["findings"]["missing_starters"], ["New Person"])

    def test_cancel_closes_the_pay_run(self):
        item = self.open_run()
        self.submit(CHECK, "Pay run", item["id"], {"q_action": "Cancel this pay run"})
        self.assertEqual((self.run_record()["status"], self.w.items[item["id"]]["status"]), ("cancelled", "done"))

    def test_nothing_is_ever_written_to_xero_payroll(self):
        item = self.open_run()
        self.check(item)
        self.submit(PACK, "Pay run", item["id"], {"q_action": "Ready for review: build the review pack"})
        self.assertTrue(all(t.startswith(("list_", "get_")) for t, _ in self.w.payroll.calls))

    # ── review pack and approvals ──
    def test_review_pack_goes_to_the_internal_reviewer(self):
        item = self.open_run()
        self.check(item)
        steps = self.submit(PACK, "Pay run", item["id"], {"q_action": "Ready for review: build the review pack",
                                                         "q_notes": "Fran's timesheet deleted; New Person added."})
        self.assertEqual(steps["build_pack"]["outputs"]["outcome"], "in_review")
        review = next(i for i in self.w.items.values() if i["title"].startswith("Review pay run: Acme Plumbing"))
        self.assertEqual((review["assigneeIds"], review["formId"]), (["user-sam"], source_id(PACK, "Payroll review")))
        pack = next(a for a in self.w.artifacts if a["title"].startswith("Pay run review pack"))
        self.assertIn("Fran&#x27;s timesheet deleted", pack["content"])
        self.assertIn("Sam Special", pack["content"])
        self.assertIn("2,560.00", pack["content"])
        # Links open in a new tab: the pack shows in the Reports viewer's frame, where a plain link loads inside the frame.
        self.assertIn("<base target='_blank'>", pack["content"])
        self.assertIn("' target='_blank' rel='noopener noreferrer'>Full check report</a>", pack["content"])
        self.assertEqual(self.run_record()["status"], "in_review")

    def test_no_draft_pay_run_means_no_pack(self):
        self.w.payroll.org(T_ACME)["pay_runs"] = []
        item = self.open_run()
        steps = self.submit(PACK, "Pay run", item["id"], {"q_action": "Ready for review: build the review pack"})
        self.assertEqual(steps["build_pack"]["outputs"]["outcome"], "no_draft")
        self.assertIn("no draft pay run", self.w.items[item["id"]]["comments"][-1])

    def approve_internally(self, item):
        self.check(item)
        self.submit(PACK, "Pay run", item["id"], {"q_action": "Ready for review: build the review pack"})
        review = next(i for i in self.w.items.values() if i["title"].startswith("Review pay run:"))
        self.submit(REVIEW, "Payroll review", review["id"], {"q_decision": "Approve: send it to the client"}, actor=SAM["id"])
        return review

    def test_internal_approval_then_client_approval_completes_the_pay_run(self):
        item = self.open_run()
        review = self.approve_internally(item)
        self.assertEqual(self.w.items[review["id"]]["status"], "done")
        client_item = next(i for i in self.w.items.values() if i["title"].startswith("Client approval: Acme Plumbing"))
        self.assertEqual(client_item["assigneeIds"], ["user-officer"])
        self.assertIn("Robin and Jordan", client_item["description"])
        self.assertEqual(self.run_record()["status"], "client_approval")
        self.submit(REVIEW, "Client payroll approval", client_item["id"], {"q_client_decision": "They approved it",
                                                                         "q_client_reply": "Looks good, go ahead"})
        self.assertEqual(self.run_record()["status"], "done")
        self.assertEqual(self.w.items[item["id"]]["status"], "done")
        self.assertIn("They post it in Xero Payroll", self.w.items[item["id"]]["comments"][-1])

    def test_when_the_practice_posts_the_item_stays_open_to_post(self):
        ps = json.loads(self.w.files["Bookkeeping/Payroll/acme-plumbing.json"])
        ps["who_posts"] = "practice"
        self.w.files["Bookkeeping/Payroll/acme-plumbing.json"] = json.dumps(ps)
        item = self.open_run()
        self.approve_internally(item)
        client_item = next(i for i in self.w.items.values() if i["title"].startswith("Client approval:"))
        self.submit(REVIEW, "Client payroll approval", client_item["id"], {"q_client_decision": "They approved it"})
        self.assertEqual(self.run_record()["status"], "approved")
        self.assertNotEqual(self.w.items[item["id"]]["status"], "done")
        self.assertIn("Post it in Xero Payroll", self.w.items[item["id"]]["comments"][-1])

    def test_changes_go_back_to_the_payroll_officer(self):
        item = self.open_run()
        self.check(item)
        self.submit(PACK, "Pay run", item["id"], {"q_action": "Ready for review: build the review pack"})
        review = next(i for i in self.w.items.values() if i["title"].startswith("Review pay run:"))
        self.submit(REVIEW, "Payroll review", review["id"], {"q_decision": "Changes needed", "q_notes": "Casey's Saturday is at ordinary rate"})
        self.assertEqual(self.run_record()["status"], "changes")
        self.assertEqual(self.w.items[item["id"]]["status"], "in_progress")
        self.assertIn("Casey's Saturday", self.w.items[item["id"]]["comments"][-1])

    # ── The checks catalogue's payroll checks ──
    def complete(self, e, bank="062-000 1111 2222", tfn="123456789"):
        """A full Xero Payroll record: tax declaration, super fund, bank account, address."""
        e["TaxDeclaration"] = dict(e.get("TaxDeclaration") or {}, TaxFileNumber=tfn)
        e["SuperMemberships"] = [{"SuperMembershipID": "SM-" + e["EmployeeID"], "SuperFundID": "F1", "EmployeeNumber": "1"}]
        e["BankAccounts"] = [{"BSB": bank.split(" ")[0], "AccountNumber": "".join(bank.split(" ")[1:]), "Remainder": True}]
        e["HomeAddress"] = {"AddressLine1": "1 High St", "City": "Sydney", "Region": "NSW", "PostalCode": "2000"}
        e["StartDate"] = xero_ms_date(START - datetime.timedelta(days=400))
        return e

    def seed_catalogue(self):
        o = self.w.payroll.org(T_ACME)
        today = datetime.datetime.now(SYD).date()
        o["leave_types"] = [{"LeaveTypeID": "AL", "Name": "Annual Leave", "LeaveCategoryCode": "ANNUALLEAVE"},
                            {"LeaveTypeID": "PL", "Name": "Personal/Carer's Leave", "LeaveCategoryCode": "PERSONALSICKCARERSLEAVE"},
                            {"LeaveTypeID": "LSL", "Name": "Long Service Leave", "LeaveCategoryCode": "LONGSERVICELEAVE"}]
        for er in o["earnings"]:
            er.update(EarningsType="ORDINARYTIMEEARNINGS", TypeOfUnits="Hours")
        o["earnings"] += [{"EarningsRateID": "OT", "Name": "Overtime x1.5", "RatePerUnit": 45.0, "EarningsType": "ORDINARYTIMEEARNINGS",
                           "TypeOfUnits": "Hours"},
                          {"EarningsRateID": "TOOL", "Name": "Tool allowance", "RatePerUnit": 10.0, "EarningsType": "ALLOWANCE",
                           "TypeOfUnits": "Hours"}]
        casey, fran, sam = o["employees"]
        for e in (casey, sam):
            self.complete(e)
        self.complete(fran, bank="062-000 1111 2222")  # the same account as Sam's (and Casey's: CHK-162)
        casey["BankAccounts"] = [{"BSB": "062-001", "AccountNumber": "99990000"}]
        # Casey (casual, 18): accrues annual leave, and has a tool allowance with no allowance type.
        casey["PayTemplate"]["LeaveLines"] = [{"LeaveTypeID": "AL", "CalculationType": "FIXEDAMOUNTEACHPERIOD", "AnnualNumberOfUnits": 152}]
        casey["PayTemplate"]["EarningsLines"].append({"EarningsRateID": "TOOL", "RatePerUnit": 10.0})
        # Fran (full-time, adult): a salary under the minimum wage, 120 hours of annual leave a year, no long service leave,
        # -8 hours of annual leave, a TFN applied for 60 days ago, no super fund.
        fran["PayTemplate"] = {"EarningsLines": [{"EarningsRateID": "ORD", "AnnualSalary": 52000.0, "NumberOfUnitsPerWeek": 38}],
                               "LeaveLines": [{"LeaveTypeID": "AL", "CalculationType": "FIXEDAMOUNTEACHPERIOD", "AnnualNumberOfUnits": 120},
                                              {"LeaveTypeID": "PL", "CalculationType": "FIXEDAMOUNTEACHPERIOD", "AnnualNumberOfUnits": 76}]}
        fran["LeaveBalances"] = [{"LeaveName": "Annual Leave", "LeaveTypeID": "AL", "NumberOfUnits": -8.0, "TypeOfUnits": "Hours"}]
        fran["TaxDeclaration"] = {"EmploymentBasis": "FULLTIME", "TFNExemptionType": "PENDING"}
        fran["StartDate"] = xero_ms_date(today - datetime.timedelta(days=60))
        fran["SuperMemberships"] = []
        # Sam (casual): overtime on a pay item set up as ordinary time.
        sam["PayTemplate"]["EarningsLines"].append({"EarningsRateID": "OT", "RatePerUnit": 45.0})
        # Terry: terminated before the period with no cessation reason, still on the draft. Nova: starts after the period,
        # full-time with no leave at all, reimbursed more than paid.
        terry = self.complete({"EmployeeID": "E4", "FirstName": "Terry", "LastName": "Gone", "Status": "TERMINATED",
                               "PayrollCalendarID": "CAL1", "DateOfBirth": xero_ms_date(datetime.date(1980, 2, 2)),
                               "TaxDeclaration": {"EmploymentBasis": "PARTTIME"}, "OrdinaryEarningsRateID": "ORD",
                               "TerminationDate": xero_ms_date(START - datetime.timedelta(days=10))})
        nova = self.complete({"EmployeeID": "E5", "FirstName": "Nova", "LastName": "Newstart", "Status": "ACTIVE",
                              "PayrollCalendarID": "CAL1", "DateOfBirth": xero_ms_date(datetime.date(1995, 3, 3)),
                              "TaxDeclaration": {"EmploymentBasis": "FULLTIME"}, "OrdinaryEarningsRateID": "ORD"},
                             bank="062-000 5555 6666")
        nova["StartDate"] = xero_ms_date(END + datetime.timedelta(days=5))
        o["employees"] += [terry, nova]
        run = o["pay_runs"][0]
        for s in run["Payslips"]:
            s["PayslipID"] = "PS-" + s["EmployeeID"]
        run["Payslips"] += [{"EmployeeID": "E4", "PayslipID": "PS-E4", "FirstName": "Terry", "LastName": "Gone", "Wages": 500,
                             "Tax": 50, "Super": 60, "Reimbursements": 0, "NetPay": 450},
                            {"EmployeeID": "E5", "PayslipID": "PS-E5", "FirstName": "Nova", "LastName": "Newstart", "Wages": 300,
                             "Tax": 0, "Super": 36, "Reimbursements": 400, "NetPay": 700}]
        # Fran's payslip: $2,000 ordinary time, $200 super guarantee (12% is $240).
        o["payslips"]["PS-E2"] = {"EmployeeID": "E2", "PayslipID": "PS-E2", "Wages": 2000, "Super": 200,
                                  "EarningsLines": [{"EarningsRateID": "ORD", "RatePerUnit": 26.32, "NumberOfUnits": 76, "Amount": 2000.0}],
                                  "SuperannuationLines": [{"ContributionType": "SGC", "CalculationType": "STATUTORY", "Amount": 200.0}]}
        # Super owed: $1,500 on the liability account; $400 of it is from a pay run 3 days ago, not due yet.
        o["accounts"] = [{"AccountID": "acc-826", "Type": "SUPERANNUATIONLIABILITY", "Code": "826", "Name": "Superannuation Payable"}]
        o["pay_runs"].append({"PayRunID": "PR0", "PayrollCalendarID": "CAL1", "PayRunStatus": "POSTED",
                              "PaymentDate": xero_ms_date(today - datetime.timedelta(days=3)), "Super": 400.0, "Payslips": []})
        self.w.xero.orgs[T_ACME]["trial_balance"] = {str(today): {"acc-826": -1500.0}}
        # Xero contacts: a supplier with Fran and Sam's account (CHK-019), and Casey's own contact for reimbursements (CHK-152).
        self.w.xero.orgs[T_ACME]["contacts"] += [
            {"ContactID": "c-sparkle", "Name": "Sparkle Supplies", "ContactStatus": "ACTIVE", "BankAccountDetails": "062000 11112222"},
            {"ContactID": "c-casey", "Name": "Casey Casual", "ContactStatus": "ACTIVE", "BankAccountDetails": "062-001 9999 0000"}]

    def test_the_catalogue_payroll_checks_find_each_problem(self):
        self.seed_catalogue()
        item = self.open_run()
        steps = self.check(item)
        self.assertEqual(steps["check"]["outputs"]["problem"], "")
        found = {}
        for r in steps["check"]["outputs"]["findings"]["catalogue"]:
            found.setdefault(r["check"], []).append(r)
        self.assertEqual(sorted(found), ["CHK-019", "CHK-147", "CHK-148", "CHK-149", "CHK-150", "CHK-151", "CHK-152", "CHK-153",
                                         "CHK-155", "CHK-156", "CHK-157", "CHK-159", "CHK-161", "CHK-162", "CHK-164", "CHK-166",
                                         "CHK-167", "CHK-168"])
        one = {c: rows[0] for c, rows in found.items()}
        self.assertEqual(one["CHK-157"]["amount"], 1100.0)
        self.assertIn("$1,100.00 of super on 826 Superannuation Payable is past its due date", one["CHK-157"]["text"])
        self.assertEqual((one["CHK-147"]["name"], one["CHK-155"]["name"], one["CHK-155"]["amount"]), ("Terry Gone", "Terry Gone", 500.0))
        self.assertEqual((one["CHK-153"]["name"], one["CHK-151"]["name"]), ("Nova Newstart", "Nova Newstart"))
        self.assertIn("is paid $26.32 an hour, below the national minimum wage of $26.44", one["CHK-156"]["text"])
        self.assertEqual(one["CHK-156"]["amount"], round((26.44 - 52000 / 52 / 38) * 76, 2))
        self.assertIn("$200.00 super guarantee on this pay run, but 12% of their ordinary time earnings ($2,000.00) is $240.00",
                      one["CHK-164"]["text"])
        self.assertIn("negative Annual Leave balance (-8.0 hours)", one["CHK-149"]["text"])
        self.assertIn("still has no tax file number 60 days after starting", one["CHK-148"]["text"])
        self.assertEqual(one["CHK-150"]["name"], "Casey Casual")
        self.assertIn("accrues 120.0 hours of annual leave a year; 4 weeks at 38.0 hours a week is 152.0", one["CHK-166"]["text"])
        self.assertEqual(sorted(r["name"] for r in found["CHK-167"]), ["Fran Fulltime", "Nova Newstart"])
        self.assertEqual([r["name"] for r in found["CHK-161"]], ["Nova Newstart"])
        self.assertEqual(one["CHK-162"]["name"], "Fran Fulltime, Sam Special")
        self.assertEqual(sorted(r["name"] for r in found["CHK-019"]), ["Fran Fulltime", "Sam Special"])
        self.assertIn("has the same bank account as the Xero contact Sparkle Supplies", one["CHK-019"]["text"])
        self.assertEqual([r["name"] for r in found["CHK-152"]], ["Casey Casual"])
        self.assertIn("Fran Fulltime: no super fund", "Fran Fulltime: " + one["CHK-159"]["text"].split("has ")[1])
        self.assertEqual(sorted(r["text"].split(" ")[2] for r in found["CHK-168"]), ["Overtime", "Tool"])
        # Severity order, the covered catalogue IDs and the heuristic flag ride along.
        self.assertEqual(steps["check"]["outputs"]["findings"]["catalogue"][0]["severity"], "High")
        self.assertEqual((one["CHK-149"]["also"], one["CHK-157"]["also"], one["CHK-164"]["also"]), (["CHK-160"], ["CHK-158"], ["CHK-165"]))
        self.assertTrue(one["CHK-156"]["heuristic"])
        # The report and the comment carry them; no tax file number, bank account or date of birth anywhere.
        comment = self.w.items[item["id"]]["comments"][-1]
        self.assertIn("Payroll checks: **" + str(sum(len(r) for r in found.values())) + "**. Fix them in Xero Payroll (6 High).", comment)
        report = next(a for a in self.w.artifacts if a["title"].startswith("Payroll check: Acme Plumbing"))
        self.assertIn("CHK-155 Terminated Employee Still Being Paid", report["content"])
        everything = report["content"] + comment + json.dumps(steps["check"]["outputs"]) + json.dumps(self.w.files.get(
            "Bookkeeping/PayRuns/acme-plumbing-2026-09-21.json"))
        for secret in ("123456789", "11112222", "1111 2222", "062-000", "2008-09-25", "1985-01-01", "25 Sep 2008"):
            self.assertNotIn(secret, everything, secret)

    def test_super_owed_needs_the_payroll_settings_and_the_ledger(self):
        f = self.check(self.open_run())["check"]["outputs"]["findings"]
        self.assertTrue(any(n.startswith("Super owed wasn't checked") for n in f["notes"]), f["notes"])
        self.assertNotIn("CHK-157", [r["check"] for r in f["catalogue"]])

    def test_super_is_due_7_business_days_after_pay_day_from_july_2026(self):
        lib = {}
        for name in ("util", "payroll"):
            with open(os.path.join(os.path.dirname(__file__), "..", "src", "lib", name + ".py"), encoding="utf-8") as fh:
                exec(fh.read(), lib)
        due = lib["pay_super_due"]
        self.assertEqual(due(datetime.date(2026, 6, 15)), datetime.date(2026, 7, 28))   # old rule: 28 days after the quarter
        # Payday Super: 7 business days (weekends and national public holidays skipped), plus a day for a state holiday.
        self.assertEqual(due(datetime.date(2026, 12, 20)), datetime.date(2027, 1, 4))   # Christmas, Boxing Day (Mon 28), New Year
        self.assertEqual(due(datetime.date(2026, 10, 6)), datetime.date(2026, 10, 16))
        self.assertEqual(due(datetime.date(2027, 3, 24)), datetime.date(2027, 4, 7))    # Good Friday 26 Mar, Easter Monday 29 Mar
        # 2027: Anzac Day falls on a Sunday (not moved nationally); Christmas and Boxing Day move to Mon 27 and Tue 28.
        self.assertEqual(lib["pay_national_holidays"](2027), {datetime.date(2027, 1, 1), datetime.date(2027, 1, 26),
                                                              datetime.date(2027, 3, 26), datetime.date(2027, 3, 29),
                                                              datetime.date(2027, 4, 25), datetime.date(2027, 12, 27),
                                                              datetime.date(2027, 12, 28)})
        self.assertEqual(due(datetime.date(2026, 3, 31)), datetime.date(2026, 4, 28))
        self.assertEqual(lib["pay_from"](lib["PAY_WAGE_FLOOR"], datetime.date(2026, 6, 30)), 24.95)
        self.assertEqual(lib["pay_from"](lib["PAY_WAGE_FLOOR"], datetime.date(2026, 7, 1)), 26.44)

    def test_payroll_items_are_visible_to_their_assignees_only(self):
        item = self.open_run()
        self.assertEqual(item["scope"], "me")
        self.assertEqual(item["assigneeIds"], [OFFICER["id"]])
        self.check(item)
        self.submit(PACK, "Pay run", item["id"], {"q_action": "Ready for review: build the review pack"}, actor=OFFICER["id"])
        review = next(i for i in self.w.items.values() if i["title"].startswith("Review pay run"))
        self.assertEqual((review["scope"], review["assigneeIds"]), ("me", [SAM["id"]]))


if __name__ == "__main__":
    unittest.main()
